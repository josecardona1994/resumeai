"""Resume generator — ported from D:/resume-engine/resume.py.
Returns bytes instead of writing files so FastAPI can stream responses.
"""

import base64
import copy
import io
import re
import tempfile
import urllib.request
from pathlib import Path

from jinja2 import Environment, FileSystemLoader

BASE_DIR = Path(__file__).parent
TEMPLATE_DIR = BASE_DIR / "templates"
ICONS_DIR = BASE_DIR / "icons"
ICONS_DIR.mkdir(exist_ok=True)


# ── Icon helpers ─────────────────────────────────────────────────────────────

def _get_icon_uri(slug: str, color: str) -> str:
    cache_file = ICONS_DIR / f"{slug}_{color}.b64"
    if cache_file.exists():
        return f"data:image/svg+xml;base64,{cache_file.read_text()}"

    ua = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
    urls = [
        f"https://cdn.jsdelivr.net/npm/simple-icons@latest/icons/{slug}.svg",
        f"https://cdn.simpleicons.org/{slug}/{color}",
    ]
    for url in urls:
        try:
            req = urllib.request.Request(url, headers=ua)
            with urllib.request.urlopen(req, timeout=8) as resp:
                svg_text = resp.read().decode("utf-8")
            hex_color = f"#{color}"
            if 'fill="' in svg_text:
                svg_text = re.sub(r'fill="[^"]*"', f'fill="{hex_color}"', svg_text, count=1)
            else:
                svg_text = svg_text.replace("<path ", f'<path fill="{hex_color}" ', 1)
            b64 = base64.b64encode(svg_text.encode("utf-8")).decode()
            cache_file.write_text(b64)
            return f"data:image/svg+xml;base64,{b64}"
        except Exception:
            continue
    return ""


def _embed_icons(profile: dict) -> dict:
    p = copy.deepcopy(profile)
    for key in ("skills", "non_cert_skills"):
        for skill in p.get(key, []):
            if isinstance(skill, dict) and skill.get("icon"):
                skill["icon_uri"] = _get_icon_uri(skill["icon"], skill.get("color", "444444"))
    return p


TEMPLATES = {
    "classic": "resume.html",
    "modern": "resume_modern.html",
    "minimal": "resume_minimal.html",
}


def render_html(profile: dict, template: str = "classic") -> str:
    p = _embed_icons(profile)
    env = Environment(loader=FileSystemLoader(str(TEMPLATE_DIR)))
    tmpl_file = TEMPLATES.get(template, "resume.html")
    return env.get_template(tmpl_file).render(**p)


# ── PDF generation ────────────────────────────────────────────────────────────

def generate_pdf_bytes(profile: dict, template: str = "classic") -> bytes:
    from playwright.sync_api import sync_playwright

    html = render_html(profile, template)

    with tempfile.TemporaryDirectory() as tmpdir:
        html_path = Path(tmpdir) / "resume.html"
        pdf_path = Path(tmpdir) / "resume.pdf"
        html_path.write_text(html, encoding="utf-8")

        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.goto(f"file:///{html_path.as_posix()}", wait_until="networkidle")
            page.wait_for_timeout(3000)
            page.pdf(
                path=str(pdf_path),
                format="A4",
                print_background=True,
                margin={"top": "0", "bottom": "0", "left": "0", "right": "0"},
            )
            browser.close()

        return pdf_path.read_bytes()


# ── DOCX generation ───────────────────────────────────────────────────────────

def generate_docx_bytes(profile: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    def add_heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'DDDDDD')
        pBdr.append(bottom)
        pPr.append(pBdr)

    personal = profile.get("personal", {})
    p = doc.add_paragraph()
    p.add_run(personal.get("name", "")).bold = True
    p.runs[0].font.size = Pt(22)
    p.paragraph_format.space_after = Pt(2)

    title_p = doc.add_paragraph()
    r = title_p.add_run(personal.get("title", "").upper())
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    title_p.paragraph_format.space_after = Pt(4)

    contact_parts = [personal.get("email", ""), personal.get("phone", ""), personal.get("location", "")]
    doc.add_paragraph("  |  ".join(p for p in contact_parts if p)).paragraph_format.space_after = Pt(8)

    add_heading("Profile")
    doc.add_paragraph(profile.get("profile_summary", ""))

    add_heading("Work Experience")
    for job in profile.get("experience", []):
        p = doc.add_paragraph()
        r = p.add_run(f"{job.get('period', '')}  |  ")
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        cr = p.add_run(job.get("company", ""))
        cr.bold = True
        cr.font.size = Pt(9)
        tp = doc.add_paragraph()
        tr = tp.add_run(job.get("title", ""))
        tr.italic = True
        tr.font.size = Pt(9)
        tp.paragraph_format.space_after = Pt(3)
        for bullet in job.get("bullets", []):
            bp = doc.add_paragraph(bullet, style="List Bullet")
            bp.paragraph_format.space_after = Pt(2)

    if profile.get("projects"):
        add_heading("Projects")
        for proj in profile["projects"]:
            p = doc.add_paragraph()
            p.add_run(proj.get("name", "")).bold = True
            if proj.get("github"):
                p.add_run(f"  —  {proj['github']}").font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            dp = doc.add_paragraph(proj.get("description", ""))
            dp.paragraph_format.space_after = Pt(6)

    add_heading("Education")
    for edu in profile.get("education", []):
        p = doc.add_paragraph()
        r = p.add_run(f"{edu.get('period', '')}  ")
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        r.font.size = Pt(8)
        ir = p.add_run(edu.get("institution", ""))
        ir.bold = True
        ir.font.size = Pt(8)
        dp = doc.add_paragraph(edu.get("degree", ""))
        dp.runs[0].italic = True
        dp.paragraph_format.space_after = Pt(8)

    add_heading("Skills")
    doc.add_paragraph("  •  ".join(s.get("name", s) if isinstance(s, dict) else s
                                    for s in profile.get("skills", [])))

    add_heading("Certificates")
    for cert in profile.get("certificates", []):
        doc.add_paragraph(cert, style="List Bullet")

    add_heading("Languages")
    for lang in profile.get("languages", []):
        p = doc.add_paragraph()
        p.add_run(lang.get("name", "")).bold = True
        p.add_run(f"  —  {lang.get('level', '')}").font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    if profile.get("non_cert_skills"):
        add_heading("Non-Certificate Skills")
        for skill in profile["non_cert_skills"]:
            name = skill.get("name", skill) if isinstance(skill, dict) else skill
            doc.add_paragraph(name, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
